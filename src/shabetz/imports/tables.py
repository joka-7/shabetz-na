"""Turning an uploaded spreadsheet into rows of text.

Excel workbooks and CSV files both arrive here and leave as a plain grid of
strings; deciding what each column means happens later, with the person
looking at it.
"""

from __future__ import annotations

import csv
import io
from datetime import date, datetime, time

MAX_ROWS = 5000
MAX_COLUMNS = 100
MAX_UPLOAD_BYTES = 5 * 1024 * 1024


class TableError(ValueError):
    """The file could not be read as a table. ``code`` says why."""

    def __init__(self, code: str, detail: str) -> None:
        self.code = code
        super().__init__(detail)


def read_table(data: bytes, filename: str) -> list[list[str]]:
    name = filename.lower()
    if name.endswith((".xlsx", ".xlsm")):
        rows = _read_workbook(data)
    elif name.endswith((".csv", ".tsv", ".txt")):
        rows = parse_delimited(_decode(data))
    elif name.endswith(".xls"):
        raise TableError(
            "old_excel",
            "Old .xls workbooks are not supported; save it as .xlsx or .csv first",
        )
    else:
        raise TableError("unsupported_file", "Upload an .xlsx, .csv or .txt file")
    return _trim(rows)


def _decode(data: bytes) -> str:
    """Decode a text file, allowing for Excel on a Hebrew Windows machine.

    "Save as CSV" there writes Windows-1255 rather than UTF-8, so a file that
    is not valid UTF-8 is read as that before giving up.
    """
    for encoding in ("utf-8-sig", "utf-16"):
        try:
            text = data.decode(encoding)
        except UnicodeDecodeError:
            continue
        if encoding == "utf-16" and not data.startswith((b"\xff\xfe", b"\xfe\xff")):
            continue
        return text
    return data.decode("cp1255", errors="replace")


def parse_delimited(text: str) -> list[list[str]]:
    """Split CSV, TSV or semicolon-separated text, whichever it is."""
    lines = text.splitlines()
    if not lines:
        return []
    sample = "\n".join(lines[:20])
    # Tabs come from pasting out of a spreadsheet and never appear in names,
    # so they win outright; otherwise whichever separator is most common.
    counts = {sep: sample.count(sep) for sep in ("\t", ",", ";")}
    delimiter = "\t" if counts["\t"] else max(counts, key=lambda sep: counts[sep])
    if counts[delimiter] == 0:
        return [[line] for line in lines]
    return [list(row) for row in csv.reader(io.StringIO(text), delimiter=delimiter)]


def _read_workbook(data: bytes) -> list[list[str]]:
    # Imported here so the dependency is only loaded by the one route using it.
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises a variety of zip and XML errors
        raise TableError("unreadable_file", "The workbook could not be opened") from exc
    try:
        sheet = workbook.active
        if sheet is None:
            return []
        rows: list[list[str]] = []
        for values in sheet.iter_rows(values_only=True):
            rows.append([_cell_text(value) for value in values[:MAX_COLUMNS]])
            if len(rows) > MAX_ROWS:
                break
        return rows
    finally:
        workbook.close()


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "yes" if value else ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time() == time() else value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return str(value).strip()


def _trim(rows: list[list[str]]) -> list[list[str]]:
    """Drop empty rows and trailing empty cells; enforce the size limits."""
    trimmed: list[list[str]] = []
    for row in rows:
        cells = [cell.strip() for cell in row[:MAX_COLUMNS]]
        while cells and not cells[-1]:
            cells.pop()
        if cells:
            trimmed.append(cells)
    if len(trimmed) > MAX_ROWS:
        raise TableError("too_many_rows", f"At most {MAX_ROWS} rows can be imported at once")
    return trimmed
